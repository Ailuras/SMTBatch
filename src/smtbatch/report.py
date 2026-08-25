#!/usr/bin/env python3
"""Export completed ``smtbatch run`` results to an Excel workbook.

Input is one or more ``<output>`` directories (or individual ``results.tsv`` /
legacy ``task_*.tsv`` files) produced by the batch runner.
Every row is one ``(file, solver)`` result.  Rows are pivoted so that a given
``file`` becomes a single row carrying one column group per solver, even when the
solvers come from different ``<output>`` directories; cross-solver consistency is
classified and a styled workbook is written.

The ``Results`` sheet uses ``path | filename | logic | file_size | consistency`` followed by
``<solver>_result | <solver>_time | <solver>_log`` and, with ``--load-output``,
``<solver>_output``. SMT-LIB details such as ``logic`` are read from the source
file during aggregation rather than copied through task TSVs.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from .task import (
    classify_consistency,
    iter_task_tsvs,
    read_file_details,
    recorded_result_row,
    result_file_timeout,
)

VALID_RESULTS = {"SAT", "UNSAT", "UNKNOWN"}
CHART_FORMAT_VERSION = "6"

# Lowercase batch-runner result -> (status, per-file result label).
RESULT_TO_STATUS = {
    "sat": ("SUCCESS", "SAT"),
    "unsat": ("SUCCESS", "UNSAT"),
    "unknown": ("SUCCESS", "UNKNOWN"),
    "timeout": ("TIMEOUT", "UNKNOWN"),
    "error": ("ERROR", "UNKNOWN"),
}


@dataclass
class LogEntry:
    solver: str
    log_path: Path
    file_path: Path
    status: str
    duration_sec: float | None
    exit_code: int | None
    raw_output: str
    result: str

    @property
    def success(self) -> bool:
        return self.status.upper() == "SUCCESS"

    @property
    def error(self) -> str:
        return "" if self.success else self.status.upper()


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export completed batch results to an Excel workbook.")
    parser.add_argument(
        "inputs",
        nargs="*",
        type=Path,
        help="one or more <output> directories (or results.tsv/task_*.tsv files). If omitted, scans 'results'.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("exports/result.xlsx"),
        help="Path to the output .xlsx file (default: exports/result.xlsx)",
    )
    parser.add_argument(
        "--load-output",
        action="store_true",
        help="Load raw solver output from output_path files into the <solver>_output column (slower).",
    )
    return parser.parse_args(argv)


def load_entries(task_files: list[Path], load_output: bool) -> list[LogEntry]:
    entries: list[LogEntry] = []
    for task_file in task_files:
        timeout = result_file_timeout(task_file)
        with task_file.open("r", encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle, delimiter="\t"):
                solver = (row.get("solver") or "").strip()
                file_str = (row.get("file") or "").strip()
                if not solver or not file_str:
                    continue
                file_path = Path(file_str)
                duration = _to_float(row.get("time"))
                exit_code = _to_int(row.get("code"))
                label = recorded_result_row(row, timeout)
                status, result = RESULT_TO_STATUS.get(label, ("ERROR", "UNKNOWN"))
                output_path = (row.get("output_path") or "").strip()
                raw_output = ""
                if load_output and output_path:
                    try:
                        raw_output = Path(output_path).read_text(encoding="utf-8", errors="replace").strip()
                    except OSError:
                        raw_output = ""
                entries.append(
                    LogEntry(
                        solver=solver,
                        log_path=Path(output_path) if output_path else task_file,
                        file_path=file_path,
                        status=status,
                        duration_sec=duration,
                        exit_code=exit_code,
                        raw_output=raw_output,
                        result=result,
                    )
                )
    return entries


def _to_float(value: str | None) -> float | None:
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _to_int(value: str | None) -> int | None:
    if not value:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def build_rows(
    entries: Iterable[LogEntry],
    include_output: bool = False,
) -> tuple[list[str], list[dict[str, object]]]:
    by_file: dict[str, dict[str, object]] = {}
    solvers: list[str] = []

    for entry in entries:
        file_key = str(entry.file_path)
        path_str = str(entry.file_path.parent)
        filename = entry.file_path.name
        if file_key not in by_file:
            details = read_file_details(entry.file_path)
            by_file[file_key] = {
                "path": path_str,
                "filename": filename,
                "logic": details.logic,
                "file_size": details.file_size,
            }
        record = by_file[file_key]
        if entry.solver not in solvers:
            solvers.append(entry.solver)
        solver_prefix = entry.solver
        if entry.success:
            result_value = entry.result
        elif entry.status == "TIMEOUT":
            result_value = "TIMEOUT"
        else:
            result_value = "ERROR" if entry.error == "" else entry.error
        record[f"{solver_prefix}_result"] = result_value
        record[f"{solver_prefix}_time"] = entry.duration_sec
        record[f"{solver_prefix}_log"] = str(entry.log_path)
        if include_output:
            record[f"{solver_prefix}_output"] = entry.raw_output

    solvers = sorted(set(solvers))
    fieldnames = [
        "path",
        "filename",
        "logic",
        "file_size",
        "consistency",
    ]
    for solver in solvers:
        fieldnames.extend([f"{solver}_result", f"{solver}_time", f"{solver}_log"])
        if include_output:
            fieldnames.append(f"{solver}_output")

    rows = []
    for file_key in sorted(by_file):
        record = by_file[file_key]
        for solver in solvers:
            record.setdefault(f"{solver}_result", "")
            record.setdefault(f"{solver}_time", None)
            record.setdefault(f"{solver}_log", "")
            if include_output:
                record.setdefault(f"{solver}_output", "")
        results = [str(record.get(f"{solver}_result", "")).upper().strip() for solver in solvers]
        record["consistency"] = classify_consistency(results)
        rows.append(record)

    return fieldnames, rows


def build_statistics(fieldnames: list[str], rows: list[dict[str, object]]) -> dict[str, object]:
    solvers = [field[:-7] for field in fieldnames if field.endswith("_result")]

    stats: dict[str, object] = {
        "solvers": solvers,
        "overall": {},
        "consistency": {"Consistent": 0, "Error": 0, "Hard": 0, "Conflict": 0, "Other": 0},
        "consistency_details": {},
    }

    for solver in solvers:
        result_col = f"{solver}_result"
        duration_col = f"{solver}_time"
        results = {"SAT": 0, "UNSAT": 0, "UNKNOWN": 0, "TIMEOUT": 0, "ERROR": 0}
        durations = []
        for row in rows:
            result = str(row.get(result_col, "")).upper().strip()
            if not result:
                continue
            if result in results:
                results[result] += 1
            else:
                results["ERROR"] += 1
            duration = row.get(duration_col)
            if duration is not None and isinstance(duration, (int, float)):
                durations.append(float(duration))
        stats["overall"][solver] = {
            "total": sum(results.values()),
            "SAT": results["SAT"],
            "UNSAT": results["UNSAT"],
            "UNKNOWN": results["UNKNOWN"],
            "TIMEOUT": results["TIMEOUT"],
            "ERROR": results["ERROR"],
            "avg_time": sum(durations) / len(durations) if durations else 0,
            "min_time": min(durations) if durations else 0,
            "max_time": max(durations) if durations else 0,
        }

    for row in rows:
        consistency_type = row.get("consistency", "Other")
        stats["consistency"][consistency_type] = stats["consistency"].get(consistency_type, 0) + 1
        if consistency_type not in stats["consistency_details"]:
            stats["consistency_details"][consistency_type] = {
                "count": 0,
                "examples": [],
            }
        details = stats["consistency_details"][consistency_type]
        details["count"] += 1
        if len(details["examples"]) < 5:
            details["examples"].append(
                {
                    "file": str(Path(str(row.get("path", ""))) / str(row.get("filename", ""))),
                    "results": {solver: row.get(f"{solver}_result", "") for solver in solvers},
                }
            )

    return stats


def build_metadata(input_roots: list[Path], output_path: Path) -> dict[str, str]:
    metadata = {
        "chart_format_version": CHART_FORMAT_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "results_root": "",
        "chart_path": str(output_path.resolve()),
    }
    if len(input_roots) != 1:
        return metadata
    root = input_roots[0].expanduser().resolve()
    if root.is_dir():
        metadata["results_root"] = str(root)
    return metadata


ILLEGAL_EXCEL_CHARS = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SECTION_FONT = Font(bold=True)
RESULT_FILLS = {
    "SAT": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "UNSAT": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "UNKNOWN": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "TIMEOUT": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "ERROR": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}
CONSISTENCY_FILLS = {
    "Consistent": PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid"),
    "Error": PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid"),
    "Hard": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "Conflict": PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),
    "Other": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
}


def sanitize_excel_value(value: object) -> object:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    if isinstance(value, str):
        cleaned = ILLEGAL_EXCEL_CHARS.sub("", value)
        if len(cleaned) > 32700:
            cleaned = cleaned[:32700] + "..."
        if cleaned and cleaned[0] in ("=", "+", "-", "@"):
            cleaned = "'" + cleaned
        return cleaned
    return str(value)


def styled_cell(worksheet: object, value: object, fill: object = None, font: object = None) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=sanitize_excel_value(value))
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    return cell


def append_header(worksheet: object, values: Iterable[object]) -> None:
    worksheet.append([styled_cell(worksheet, value, HEADER_FILL, HEADER_FONT) for value in values])


def append_section(worksheet: object, title: str) -> None:
    worksheet.append([styled_cell(worksheet, title, SECTION_FILL, SECTION_FONT)])


def result_fill(value: object) -> object:
    normalized = str(value).upper().strip() if value is not None else ""
    if normalized in RESULT_FILLS:
        return RESULT_FILLS[normalized]
    if normalized.startswith("ERROR") or (normalized and normalized not in {"SAT", "UNSAT", "UNKNOWN", "TIMEOUT"}):
        return RESULT_FILLS["ERROR"]
    return None


def result_column_width(field: str) -> float:
    fixed = {
        "path": 80,
        "filename": 50,
        "logic": 18,
        "file_size": 16,
        "consistency": 16,
    }
    if field in fixed:
        return fixed[field]
    if field.endswith("_result"):
        return 14
    if field.endswith("_time"):
        return 14
    if field.endswith("_log") or field.endswith("_output"):
        return 80
    return 20


def write_metadata_sheet(workbook: Workbook, metadata: dict[str, str]) -> None:
    worksheet = workbook.create_sheet("Metadata")
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 100
    append_header(worksheet, ["key", "value"])
    for key, value in metadata.items():
        worksheet.append([sanitize_excel_value(key), sanitize_excel_value(value)])


def write_excel(output_path: Path, fieldnames: list[str], rows: list[dict[str, object]], metadata: dict[str, str]) -> None:
    stats = build_statistics(fieldnames, rows)
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Results")
    worksheet.freeze_panes = "A2"
    for index, field in enumerate(fieldnames, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = result_column_width(field)
    append_header(worksheet, fieldnames)

    for row_data in rows:
        cells: list[object] = []
        for field in fieldnames:
            value = sanitize_excel_value(row_data.get(field))
            fill = result_fill(value) if field.endswith("_result") else None
            if field == "consistency":
                fill = CONSISTENCY_FILLS.get(str(value))
            cells.append(styled_cell(worksheet, value, fill=fill) if fill is not None else value)
        worksheet.append(cells)

    write_statistics_sheet(workbook, stats)
    write_metadata_sheet(workbook, metadata)
    workbook.save(output_path)


def append_solver_statistics(worksheet: object, solver: str, data: dict[str, object], include_min_max: bool) -> None:
    values = [
        solver,
        data["total"],
        data["SAT"],
        data["UNSAT"],
        data["UNKNOWN"],
        data["TIMEOUT"],
        data["ERROR"],
        round(data["avg_time"], 4),
    ]
    if include_min_max:
        values.extend([round(data["min_time"], 4), round(data["max_time"], 4)])
    cells: list[object] = []
    result_names = [None, None, "SAT", "UNSAT", "UNKNOWN", "TIMEOUT", "ERROR"]
    for index, value in enumerate(values):
        fill = RESULT_FILLS[result_names[index]] if index < len(result_names) and result_names[index] else None
        cells.append(styled_cell(worksheet, value, fill=fill) if fill is not None else value)
    worksheet.append(cells)


def write_statistics_sheet(workbook: Workbook, stats: dict[str, object]) -> None:
    worksheet = workbook.create_sheet("Statistics")
    widths = [36, 16, 14, 14, 14, 14, 14, 18, 18, 18]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    solvers = stats["solvers"]

    append_section(worksheet, "Overall Statistics")
    append_header(
        worksheet,
        ["Solver", "Total", "SAT", "UNSAT", "UNKNOWN", "TIMEOUT", "ERROR", "Avg Time (s)", "Min Time (s)", "Max Time (s)"],
    )
    for solver in solvers:
        append_solver_statistics(worksheet, solver, stats["overall"][solver], include_min_max=True)

    worksheet.append([None])
    append_section(worksheet, "Result Consistency Analysis")
    append_header(worksheet, ["Category", "Count", "Percentage", "Description"])
    total_cases = sum(stats["consistency"].values())
    descriptions = {
        "Consistent": "Simple and consistent (all solvers agree on SAT/UNSAT)",
        "Error": "Original problem has issues (all solvers failed)",
        "Hard": "Difficult problem (only UNKNOWN/TIMEOUT/ERROR, no SAT/UNSAT)",
        "Conflict": "CRITICAL: Inconsistency detected (both SAT and UNSAT)",
        "Other": "Other cases (mixed results)",
    }
    for category in ["Consistent", "Conflict", "Hard", "Error", "Other"]:
        count = stats["consistency"].get(category, 0)
        percentage = count / total_cases * 100 if total_cases else 0
        worksheet.append(
            [
                styled_cell(worksheet, category, fill=CONSISTENCY_FILLS.get(category)),
                count,
                f"{percentage:.2f}%",
                descriptions[category],
            ]
        )

def _default_results_root() -> Path:
    try:
        from .config import load_config

        return load_config().results_root
    except RuntimeError:
        return Path("results")


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    if args.inputs:
        metadata_inputs = args.inputs
    else:
        results_root = _default_results_root()
        if not results_root.exists():
            print(f"error: default results directory not found: {results_root}", file=sys.stderr)
            return 1
        metadata_inputs = [results_root]

    task_files = iter_task_tsvs(metadata_inputs)
    if not task_files:
        print("error: no complete batch results found", file=sys.stderr)
        return 1

    entries = load_entries(task_files, args.load_output)
    if not entries:
        print("warning: no result rows found", file=sys.stderr)

    fieldnames, rows = build_rows(entries, include_output=args.load_output)
    output_path = args.output.expanduser()
    if output_path.suffix.lower() != ".xlsx":
        print("error: --output must end in .xlsx", file=sys.stderr)
        return 2
    metadata = build_metadata(metadata_inputs, output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_excel(output_path, fieldnames, rows, metadata)
    print(f"[export] result_sets={len(task_files)} files={len(rows)} -> {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
