"""Plan, execute, resume, inspect, and report frozen reduction studies.

Reduction is deliberately benchmark-centric: every benchmark owns the
predicate (and therefore the solver invocation).  Reducers are the experiment
arms.  A job is one ``benchmark x reducer x repeat`` trial.
"""

from __future__ import annotations

import argparse
import concurrent.futures
from collections import Counter
import csv
import fcntl
import hashlib
import json
import math
import os
from pathlib import Path
import platform
import re
import shutil
import signal
import socket
import statistics
import subprocess
import sys
import threading
import time
from datetime import datetime, timezone
from typing import Iterable, Mapping, Sequence

from .config import Config, ReducerSpec, load_config, validate_target_branch
from . import provenance


SCHEMA_VERSION = 3
FORMAT = "reduction-v3"
SUPPORTED_FORMATS = {2: "reduction-v2", 3: FORMAT}
ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")
MAX_LOG_BYTES = 64 * 1024
STUDY_FIELDS = {
    "schema_version", "kind", "study_id", "root", "execution",
    "predicate_wrapper", "benchmarks", "reducers", "repeats",
    "limits", "comparisons", "catalog",
}
BENCHMARK_FIELDS = {
    "id", "input", "family", "theory", "predicate_mode", "solver",
    "predicate",
}
PREDICATE_FIELDS = {"command", "match"}
MATCH_FIELDS = {"ignore_stdout", "ignore_stderr", "match_stdout", "match_stderr"}
WRAPPER_FIELDS = {"command", "env"}
EXECUTION_FIELDS = {"outer_jobs", "schedule"}
LIMIT_FIELDS = {
    "trial_wall_sec", "predicate_timeout_sec", "memory_mb",
    "preflight_repeats", "verification_repeats", "termination_grace_sec",
    "analysis_horizon_sec",
}
TOOL_LIST_PLACEHOLDERS = {"{predicate}"}
TOOL_SCALAR_PLACEHOLDERS = {
    "input", "output", "workdir", "predicate_timeout",
}
WRAPPER_LIST_PLACEHOLDERS = {"{command}", "{match_args}"}
WRAPPER_SCALAR_PLACEHOLDERS = {
    "journal", "phase", "predicate_timeout", "run_id", "job_id", "attempt",
}
RESULT_FIELDS = [
    "job_id", "benchmark", "reducer", "repeat", "wave", "status",
    "status_detail", "verified", "evidence_ok", "input_expressions",
    "input_nodes", "input_bytes", "output_expressions", "output_nodes",
    "output_bytes", "size_ratio", "trial_wall_sec", "cleanup_wall_sec",
    "predicate_calls", "accepted_moves", "attempt", "output",
]
RUNNING_STATES = {"starting", "running", "stopping", "aborting", "resuming"}
FINAL_STATES = {"complete", "interrupted", "failed"}


class ReductionError(RuntimeError):
    """Report malformed studies, drift, or incomplete reduction evidence."""


class ImmediateAbort(ReductionError):
    """Stop one in-flight trial without sealing it as complete."""


def _schema_identity(value: Mapping[str, object], label: str) -> tuple[int, str]:
    version = value.get("schema_version")
    format_value = value.get("format")
    if (
        isinstance(version, bool)
        or not isinstance(version, int)
        or SUPPORTED_FORMATS.get(version) != format_value
    ):
        raise ReductionError(f"unsupported {label} schema/format")
    return version, str(format_value)


class _RunLock:
    """Cross-process run lock with a controller pid visible to the service."""

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir
        self.pid_path = output_dir / ".run.pid"
        self._handle = None

    def acquire(self) -> None:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        handle = (self.output_dir / ".run.lock").open("a+", encoding="utf-8")
        try:
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            handle.close()
            raise ValueError(f"another reduction controller is already using {self.output_dir}") from None
        try:
            handle.seek(0)
            handle.truncate()
            handle.write(f"{os.getpid()}\n")
            handle.flush()
            os.fsync(handle.fileno())
            _atomic_write(self.pid_path, f"{os.getpid()}\n".encode("utf-8"))
        except Exception:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            handle.close()
            raise
        self._handle = handle

    def release(self) -> None:
        handle = self._handle
        if handle is None:
            return
        try:
            try:
                owner = int(self.pid_path.read_text(encoding="utf-8").strip())
            except (OSError, ValueError):
                owner = None
            if owner == os.getpid():
                self.pid_path.unlink(missing_ok=True)
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
            handle.close()
            self._handle = None


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _json_bytes(value: object) -> bytes:
    return (
        json.dumps(value, sort_keys=True, ensure_ascii=True, allow_nan=False, separators=(",", ":"))
        + "\n"
    ).encode("utf-8")


def _atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with temporary.open("wb") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _write_json(path: Path, value: object) -> None:
    _atomic_write(path, _json_bytes(value))


def _append_jsonl(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            handle.write(_json_bytes(value).decode("utf-8"))
            handle.flush()
            os.fsync(handle.fileno())
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _read_json(path: Path, label: str) -> object:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ReductionError(f"unable to read {label} {path}: {exc}") from exc


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _hash_json(value: object) -> str:
    return _sha256_bytes(_json_bytes(value))


def _mapping(value: object, label: str) -> dict[str, object]:
    if not isinstance(value, dict) or not all(isinstance(key, str) for key in value):
        raise ReductionError(f"{label} must be an object with string keys")
    return dict(value)


def _only_fields(value: Mapping[str, object], allowed: set[str], label: str) -> None:
    extras = sorted(set(value) - allowed)
    if extras:
        raise ReductionError(f"unknown {label} fields: {', '.join(extras)}")


def _identifier(value: object, label: str) -> str:
    if not isinstance(value, str) or not ID_RE.fullmatch(value):
        raise ReductionError(f"{label} must match {ID_RE.pattern}")
    return value


def _string_list(value: object, label: str, *, nonempty: bool = False) -> list[str]:
    if not isinstance(value, list) or not all(isinstance(item, str) for item in value):
        raise ReductionError(f"{label} must be a list of strings")
    if nonempty and not value:
        raise ReductionError(f"{label} must not be empty")
    return list(value)


def _string_mapping(value: object, label: str) -> dict[str, str]:
    if value is None:
        return {}
    if not isinstance(value, dict) or not all(
        isinstance(key, str) and isinstance(item, str) for key, item in value.items()
    ):
        raise ReductionError(f"{label} must be an object of string values")
    return dict(value)


def _positive_int(value: object, label: str, *, allow_zero: bool = False) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise ReductionError(f"{label} must be an integer")
    if value < 0 or (value == 0 and not allow_zero):
        qualifier = "non-negative" if allow_zero else "positive"
        raise ReductionError(f"{label} must be {qualifier}")
    return value


def _positive_number(value: object, label: str, *, allow_zero: bool = False) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ReductionError(f"{label} must be a number")
    number = float(value)
    if not math.isfinite(number) or number < 0 or (number == 0 and not allow_zero):
        qualifier = "non-negative" if allow_zero else "positive"
        raise ReductionError(f"{label} must be finite and {qualifier}")
    return number


def _resolve_root(study_path: Path, value: object) -> Path:
    raw = Path(str(value if value is not None else ".")).expanduser()
    root = (study_path.parent / raw).resolve() if not raw.is_absolute() else raw.resolve()
    if not root.is_dir():
        raise ReductionError(f"study root is not a directory: {root}")
    return root


def _resolve_file(root: Path, value: object, label: str) -> Path:
    if not isinstance(value, str) or not value:
        raise ReductionError(f"{label} must be a path string")
    candidate = Path(value).expanduser()
    candidate = root / candidate if not candidate.is_absolute() else candidate
    if candidate.is_symlink():
        raise ReductionError(f"{label} must not be a symbolic link: {candidate}")
    path = candidate.resolve()
    if not path.is_file():
        raise ReductionError(f"{label} is not a regular file: {path}")
    return path


def _fingerprint(path: Path, cache: dict[Path, dict[str, object]] | None = None) -> dict[str, object]:
    resolved = path.resolve()
    if cache is not None and resolved in cache:
        return cache[resolved]
    value = {"path": str(resolved), "bytes": resolved.stat().st_size, "sha256": _sha256_path(resolved)}
    if cache is not None:
        cache[resolved] = value
    return value


def _normalize_match(value: object, label: str) -> dict[str, object]:
    raw = _mapping(value or {}, label)
    _only_fields(raw, MATCH_FIELDS, label)
    result: dict[str, object] = {}
    for key in ("ignore_stdout", "ignore_stderr"):
        item = raw.get(key, False)
        if not isinstance(item, bool):
            raise ReductionError(f"{label}.{key} must be boolean")
        result[key] = item
    for key in ("match_stdout", "match_stderr"):
        item = raw.get(key)
        if item is not None and not isinstance(item, str):
            raise ReductionError(f"{label}.{key} must be a string or null")
        result[key] = item
    return result


def _normalize_limits(value: object) -> dict[str, object]:
    raw = _mapping(value, "limits")
    _only_fields(raw, LIMIT_FIELDS, "limits")
    trial = _positive_number(raw.get("trial_wall_sec", 3600), "limits.trial_wall_sec")
    return {
        "trial_wall_sec": trial,
        "predicate_timeout_sec": _positive_number(
            raw.get("predicate_timeout_sec", 25), "limits.predicate_timeout_sec"
        ),
        "memory_mb": _positive_int(raw.get("memory_mb", 0), "limits.memory_mb", allow_zero=True),
        "preflight_repeats": _positive_int(
            raw.get("preflight_repeats", 1), "limits.preflight_repeats"
        ),
        "verification_repeats": _positive_int(
            raw.get("verification_repeats", 1), "limits.verification_repeats"
        ),
        "termination_grace_sec": _positive_number(
            raw.get("termination_grace_sec", 5), "limits.termination_grace_sec"
        ),
        "analysis_horizon_sec": _positive_number(
            raw.get("analysis_horizon_sec", trial), "limits.analysis_horizon_sec"
        ),
    }


def _validate_template(
    command: Sequence[str], *, list_placeholders: set[str], scalar_placeholders: set[str], label: str
) -> None:
    for placeholder in list_placeholders:
        if list(command).count(placeholder) > 1:
            raise ReductionError(f"{label} may contain {placeholder} at most once")
    scalar_re = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")
    allowed = {item.strip("{}") for item in list_placeholders} | scalar_placeholders
    for token in command:
        for match in scalar_re.finditer(token):
            if match.group(1) not in allowed:
                raise ReductionError(f"{label} has unknown placeholder {{{match.group(1)}}}")


def load_study(path: Path) -> dict[str, object]:
    study_path = path.expanduser().resolve()
    raw = _mapping(_read_json(study_path, "study"), "study")
    _only_fields(raw, STUDY_FIELDS, "study")
    if raw.get("schema_version") != SCHEMA_VERSION or raw.get("kind") != "reduction":
        raise ReductionError("study must use schema_version 3 and kind 'reduction'")
    study_id = _identifier(raw.get("study_id"), "study_id")
    root = _resolve_root(study_path, raw.get("root"))
    fingerprints: dict[Path, dict[str, object]] = {}

    execution_raw = _mapping(raw.get("execution"), "execution")
    _only_fields(execution_raw, EXECUTION_FIELDS, "execution")
    schedule = execution_raw.get("schedule", "strict-wave")
    if schedule != "strict-wave":
        raise ReductionError("execution.schedule must be 'strict-wave'")
    execution = {
        "outer_jobs": _positive_int(execution_raw.get("outer_jobs"), "execution.outer_jobs"),
        "schedule": schedule,
    }

    wrapper_raw = _mapping(raw.get("predicate_wrapper"), "predicate_wrapper")
    _only_fields(wrapper_raw, WRAPPER_FIELDS, "predicate_wrapper")
    wrapper_command = _string_list(
        wrapper_raw.get("command"), "predicate_wrapper.command", nonempty=True
    )
    _validate_template(
        wrapper_command,
        list_placeholders=WRAPPER_LIST_PLACEHOLDERS,
        scalar_placeholders=WRAPPER_SCALAR_PLACEHOLDERS,
        label="predicate_wrapper.command",
    )
    for required in ("{command}", "{journal}", "{phase}"):
        if required == "{command}":
            present = wrapper_command.count(required) == 1
        else:
            present = any(required in token for token in wrapper_command)
        if not present:
            raise ReductionError(f"predicate_wrapper.command must contain {required}")
    wrapper = {
        "command": wrapper_command,
        "env": _string_mapping(wrapper_raw.get("env"), "predicate_wrapper.env"),
    }

    benchmarks_raw = raw.get("benchmarks")
    if not isinstance(benchmarks_raw, list) or not benchmarks_raw:
        raise ReductionError("benchmarks must be a non-empty list")
    benchmarks = []
    benchmark_ids: set[str] = set()
    for index, item in enumerate(benchmarks_raw):
        benchmark = _mapping(item, f"benchmarks[{index}]")
        _only_fields(benchmark, BENCHMARK_FIELDS, f"benchmarks[{index}]")
        benchmark_id = _identifier(benchmark.get("id"), f"benchmarks[{index}].id")
        if benchmark_id in benchmark_ids:
            raise ReductionError(f"duplicate benchmark id: {benchmark_id}")
        benchmark_ids.add(benchmark_id)
        for field in ("family", "theory", "predicate_mode", "solver"):
            if field not in benchmark:
                raise ReductionError(f"benchmark {benchmark_id} is missing {field}")
        for field in ("family", "theory", "predicate_mode"):
            if not isinstance(benchmark[field], str):
                raise ReductionError(f"benchmark {benchmark_id}.{field} must be a string")
        solver_metadata = _mapping(
            benchmark.get("solver"), f"benchmark {benchmark_id} solver"
        )
        input_path = _resolve_file(root, benchmark.get("input"), f"benchmark {benchmark_id} input")
        input_fingerprint = _fingerprint(input_path, fingerprints)
        predicate = _mapping(benchmark.get("predicate"), f"benchmark {benchmark_id} predicate")
        _only_fields(predicate, PREDICATE_FIELDS, f"benchmark {benchmark_id} predicate")
        predicate_command = _string_list(
            predicate.get("command"), f"benchmark {benchmark_id} predicate.command", nonempty=True
        )
        benchmarks.append({
            "id": benchmark_id,
            "input": str(input_path),
            "input_bytes": input_fingerprint["bytes"],
            "input_sha256": input_fingerprint["sha256"],
            "family": benchmark["family"],
            "theory": benchmark["theory"],
            "predicate_mode": benchmark["predicate_mode"],
            "solver": solver_metadata,
            "predicate": {
                "command": predicate_command,
                "match": _normalize_match(
                    predicate.get("match"), f"benchmark {benchmark_id} predicate.match"
                ),
            },
        })

    reducers_raw = raw.get("reducers")
    reducers = _string_list(reducers_raw, "reducers", nonempty=True)
    reducer_ids: set[str] = set()
    for index, reducer in enumerate(reducers):
        reducer_id = _identifier(reducer, f"reducers[{index}]")
        if reducer_id in reducer_ids:
            raise ReductionError(f"duplicate reducer id: {reducer_id}")
        reducer_ids.add(reducer_id)

    repeats = _positive_int(raw.get("repeats", 1), "repeats")
    comparisons_raw = raw.get("comparisons", [])
    if not isinstance(comparisons_raw, list):
        raise ReductionError("comparisons must be a list")
    comparisons = []
    for index, pair in enumerate(comparisons_raw):
        if (
            not isinstance(pair, list) or len(pair) != 2
            or not all(isinstance(item, str) for item in pair)
            or pair[0] == pair[1] or any(item not in reducer_ids for item in pair)
        ):
            raise ReductionError(f"comparisons[{index}] must reference two different reducers")
        if pair not in comparisons:
            comparisons.append(pair)

    catalog = raw.get("catalog", {})
    if not isinstance(catalog, dict):
        raise ReductionError("catalog must be an object")

    return {
        "schema_version": SCHEMA_VERSION,
        "kind": "reduction",
        "study_id": study_id,
        "root": str(root),
        "execution": execution,
        "predicate_wrapper": wrapper,
        "benchmarks": benchmarks,
        "reducers": reducers,
        "repeats": repeats,
        "limits": _normalize_limits(raw.get("limits", {})),
        "comparisons": comparisons,
        "catalog": dict(catalog),
        "source": {"path": str(study_path), "sha256": _sha256_path(study_path)},
        "environment": {
            "python": sys.version,
            "python_executable": sys.executable,
            "platform": platform.platform(),
            "hostname": socket.gethostname(),
            "cpu_count": os.cpu_count(),
        },
    }


def _stable_offset(study_id: str, benchmark_id: str, repeat: int, reducer_count: int) -> int:
    identity = f"{study_id}\0{benchmark_id}\0{repeat}".encode("utf-8")
    return int(hashlib.sha256(identity).hexdigest()[:16], 16) % reducer_count


def _plan_reducer(spec: ReducerSpec, root: Path) -> dict[str, object]:
    command = list(spec.command)
    _validate_template(
        command,
        list_placeholders=TOOL_LIST_PLACEHOLDERS,
        scalar_placeholders=TOOL_SCALAR_PLACEHOLDERS,
        label=f"reducer {spec.name} command",
    )
    for required in ("{input}", "{output}", "{predicate}"):
        if not any(required in token for token in command):
            raise ReductionError(f"reducer {spec.name} command must contain {required}")
    try:
        frozen_provenance = provenance.snapshot_reducer(
            spec.executable, spec.provenance_paths, require_clean=spec.require_clean
        )
    except provenance.ProvenanceError as exc:
        raise ReductionError(f"unable to freeze reducer {spec.name}: {exc}") from exc
    return {
        "id": spec.name,
        "label": spec.label,
        "command": command,
        "env": dict(spec.env),
        "provenance": frozen_provenance,
    }


def build_plan(
    study: Mapping[str, object], *, config: Config | None = None,
    reducers: Sequence[str] | None = None, timeout_seconds: float | None = None,
    outer_jobs: int | None = None, benchmark_ids: Sequence[str] | None = None,
    repeats: int | None = None, selection: Mapping[str, object] | None = None,
) -> dict[str, object]:
    root = Path(str(study["root"]))
    try:
        config = config or load_config(root)
    except RuntimeError as exc:
        raise ReductionError(str(exc)) from None
    allowed = [str(item) for item in study["reducers"]]
    selected_ids = list(reducers) if reducers is not None else list(allowed)
    if not selected_ids:
        raise ReductionError("at least one reducer must be selected")
    if len(set(selected_ids)) != len(selected_ids):
        raise ReductionError("selected reducers must not contain duplicates")
    unknown = [item for item in selected_ids if item not in config.reducers]
    if unknown:
        raise ReductionError(f"unknown configured reducer: {', '.join(unknown)}")
    forbidden = [item for item in selected_ids if item not in allowed]
    if forbidden:
        raise ReductionError(f"reducers not allowed by study: {', '.join(forbidden)}")
    resolved_reducers = [_plan_reducer(config.reducers[item], root) for item in selected_ids]
    trial_timeout = (
        _positive_number(timeout_seconds, "timeout_seconds")
        if timeout_seconds is not None else float(study["limits"]["trial_wall_sec"])
    )
    workers = (
        _positive_int(outer_jobs, "outer_jobs")
        if outer_jobs is not None else int(study["execution"]["outer_jobs"])
    )
    limits = dict(study["limits"])
    limits["trial_wall_sec"] = trial_timeout
    execution = {"outer_jobs": workers, "schedule": "strict-wave"}
    configured_comparisons = list(config.comparisons) or list(study["comparisons"])
    comparisons = [
        list(pair) for pair in configured_comparisons
        if pair[0] in selected_ids and pair[1] in selected_ids
    ]
    reducers = resolved_reducers
    all_benchmarks = list(study["benchmarks"])
    by_benchmark = {str(item["id"]): item for item in all_benchmarks}
    if benchmark_ids is None:
        selected_benchmark_ids = list(by_benchmark)
    else:
        selected_benchmark_ids = list(benchmark_ids)
        if not selected_benchmark_ids:
            raise ReductionError("at least one benchmark must be selected")
        if len(set(selected_benchmark_ids)) != len(selected_benchmark_ids):
            raise ReductionError("selected benchmarks must not contain duplicates")
        unknown_benchmarks = [item for item in selected_benchmark_ids if item not in by_benchmark]
        if unknown_benchmarks:
            raise ReductionError(
                f"unknown benchmark: {', '.join(unknown_benchmarks)}"
            )
    benchmarks = [by_benchmark[item] for item in selected_benchmark_ids]
    repeat_count = (
        _positive_int(repeats, "repeats") if repeats is not None else int(study["repeats"])
    )
    selection_record = dict(selection or {})
    jobs = []
    order = 0
    reducer_count = len(reducers)
    for repeat in range(1, repeat_count + 1):
        offsets = {
            benchmark["id"]: _stable_offset(
                str(study["study_id"]), str(benchmark["id"]), repeat, reducer_count
            )
            for benchmark in benchmarks
        }
        for slot in range(reducer_count):
            wave = (repeat - 1) * reducer_count + slot + 1
            for benchmark_index, benchmark in enumerate(benchmarks, start=1):
                reducer = reducers[(offsets[benchmark["id"]] + slot) % reducer_count]
                order += 1
                identity = f"{study['study_id']}\0{benchmark['id']}\0{reducer['id']}\0{repeat}"
                jobs.append({
                    "job_id": hashlib.sha256(identity.encode("utf-8")).hexdigest()[:20],
                    "order": order,
                    "wave": wave,
                    "slot": slot + 1,
                    "benchmark_index": benchmark_index,
                    "benchmark_id": benchmark["id"],
                    "reducer_id": reducer["id"],
                    "repeat": repeat,
                })
    try:
        harness_provenance = provenance.snapshot_command(
            list(study["predicate_wrapper"]["command"]), cwd=root, execute=False
        )
    except provenance.ProvenanceError as exc:
        raise ReductionError(f"unable to freeze predicate wrapper: {exc}") from exc
    plan = {
        "schema_version": SCHEMA_VERSION,
        "format": FORMAT,
        "study_id": study["study_id"],
        "created_at": _utc_now(),
        "root": study["root"],
        "execution": execution,
        "predicate_wrapper": study["predicate_wrapper"],
        "harness_provenance": harness_provenance,
        "benchmarks": benchmarks,
        "reducers": reducers,
        "repeats": repeat_count,
        "limits": limits,
        "comparisons": comparisons,
        "catalog": study.get("catalog", {}),
        "selection": selection_record,
        "source": study["source"],
        "environment": study["environment"],
        "jobs": jobs,
        "job_count": len(jobs),
        "wave_count": repeat_count * reducer_count,
    }
    plan["plan_sha256"] = _hash_json(plan)
    return plan


def _write_jobs(path: Path, jobs: Sequence[Mapping[str, object]]) -> None:
    lines = ["job_id\torder\twave\tbenchmark\treducer\trepeat\n"]
    for job in jobs:
        lines.append(
            f"{job['job_id']}\t{job['order']}\t{job['wave']}\t{job['benchmark_id']}\t"
            f"{job['reducer_id']}\t{job['repeat']}\n"
        )
    _atomic_write(path, "".join(lines).encode("utf-8"))


def _validate_frozen_file(value: object, label: str) -> None:
    record = _mapping(value, label)
    path_value = record.get("path")
    expected = record.get("sha256")
    if not isinstance(path_value, str) or not isinstance(expected, str):
        raise ReductionError(f"malformed {label} fingerprint")
    path = Path(path_value)
    if path.is_symlink() or not path.is_file():
        raise ReductionError(f"{label} is missing or not a regular file: {path}")
    if _sha256_path(path) != expected:
        raise ReductionError(f"{label} drift: {path}")


def _validate_live_provenance(plan: Mapping[str, object]) -> None:
    """Reject any live asset that differs from a freshly prepared v3 plan."""

    version, _ = _schema_identity(plan, "plan")
    if version != SCHEMA_VERSION:
        raise ReductionError(
            "reduction-v2 plans are read-only; prepare a reduction-v3 plan before execution"
        )

    source = _mapping(plan.get("source"), "plan source")
    _validate_frozen_file(source, "study source")
    for benchmark in plan.get("benchmarks", []):
        item = _mapping(benchmark, "benchmark")
        _validate_frozen_file(
            {"path": item.get("input"), "sha256": item.get("input_sha256")},
            f"benchmark {item.get('id')} input",
        )

    reducers = plan.get("reducers")
    if not isinstance(reducers, list):
        raise ReductionError("plan reducers must be a list")
    for item in reducers:
        reducer = _mapping(item, "plan reducer")
        frozen = _mapping(
            reducer.get("provenance"), f"reducer {reducer.get('id')} provenance"
        )
        try:
            current = provenance.resnapshot_reducer(frozen)
        except provenance.ProvenanceError as exc:
            raise ReductionError(
                f"reducer {reducer.get('id')} provenance drift: {exc}"
            ) from exc
        if not provenance.same_snapshot(frozen, current):
            raise ReductionError(f"reducer {reducer.get('id')} provenance drift")

    frozen_harness = _mapping(plan.get("harness_provenance"), "predicate wrapper provenance")
    try:
        current_harness = provenance.resnapshot_command(frozen_harness)
    except provenance.ProvenanceError as exc:
        raise ReductionError(f"predicate wrapper provenance drift: {exc}") from exc
    if not provenance.same_snapshot(frozen_harness, current_harness):
        raise ReductionError("predicate wrapper provenance drift")

    catalog = plan.get("catalog", {})
    if not isinstance(catalog, dict):
        raise ReductionError("plan catalog must be an object")
    for key in ("database", "template"):
        record = catalog.get(key)
        if record is not None:
            _validate_frozen_file(record, f"benchmark catalog {key}")
    frozen_identity = catalog.get("identity")
    if frozen_identity is not None:
        identity_record = _mapping(frozen_identity, "benchmark identity")
        try:
            current_identity = provenance.resnapshot_command(identity_record)
        except provenance.ProvenanceError as exc:
            raise ReductionError(f"benchmark identity drift: {exc}") from exc
        if not provenance.same_snapshot(identity_record, current_identity):
            raise ReductionError("benchmark identity drift")


def prepare(
    study_path: Path, output: Path, *, reducers: Sequence[str] | None = None,
    timeout_seconds: float | None = None, outer_jobs: int | None = None,
    benchmark_ids: Sequence[str] | None = None, repeats: int | None = None,
    selection: Mapping[str, object] | None = None,
) -> dict[str, object]:
    study = load_study(study_path)
    try:
        project_config = load_config(Path(str(study["root"])))
        validate_target_branch(project_config)
    except RuntimeError as exc:
        raise ReductionError(str(exc)) from None
    plan = build_plan(
        study, config=project_config, reducers=reducers,
        timeout_seconds=timeout_seconds, outer_jobs=outer_jobs,
        benchmark_ids=benchmark_ids, repeats=repeats, selection=selection,
    )
    _validate_live_provenance(plan)
    output = output.expanduser().resolve()
    if output.exists() and any(output.iterdir()):
        existing = output / "plan.json"
        if existing.is_file():
            loaded = load_plan(output)
            comparable = ("source", "execution", "reducers", "limits", "repeats", "selection", "jobs")
            if all(loaded.get(key) == plan.get(key) for key in comparable):
                return loaded
        raise ReductionError(f"refusing to prepare a non-empty output directory: {output}")
    output.mkdir(parents=True, exist_ok=True)
    _write_json(output / "study.json", study)
    _write_json(output / "plan.json", plan)
    _write_jobs(output / "jobs.tsv", plan["jobs"])
    _write_json(
        output / "provenance.json",
        {
            "schema_version": SCHEMA_VERSION,
            "format": FORMAT,
            "study_source": study["source"],
            "environment": study["environment"],
            "reducers": {
                item["id"]: item["provenance"] for item in plan["reducers"]
            },
            "predicate_wrapper": plan["harness_provenance"],
            "catalog": plan.get("catalog", {}),
            "inputs": [
                {"id": item["id"], "path": item["input"], "sha256": item["input_sha256"]}
                for item in plan["benchmarks"]
            ],
        },
    )
    _write_json(
        output / "plan.complete.json",
        {
            "schema_version": SCHEMA_VERSION,
            "format": FORMAT,
            "plan_sha256": _sha256_path(output / "plan.json"),
            "jobs_sha256": _sha256_path(output / "jobs.tsv"),
            "study_sha256": _sha256_path(output / "study.json"),
            "provenance_sha256": _sha256_path(output / "provenance.json"),
        },
    )
    _write_json(
        output / "progress.json",
        {
            "schema_version": SCHEMA_VERSION,
            "format": FORMAT,
            "study_id": plan["study_id"],
            "status": "prepared",
            "updated_at": _utc_now(),
            "total_jobs": len(plan["jobs"]),
            "completed_jobs": 0,
            "running_jobs": 0,
            "pending_jobs": len(plan["jobs"]),
        },
    )
    return plan


def load_plan(output: Path) -> dict[str, object]:
    output = output.expanduser().resolve()
    plan = _mapping(_read_json(output / "plan.json", "plan"), "plan")
    marker = _mapping(_read_json(output / "plan.complete.json", "plan marker"), "plan marker")
    plan_version, plan_format = _schema_identity(plan, "plan")
    marker_version, marker_format = _schema_identity(marker, "plan marker")
    if (
        marker_version != plan_version
        or marker_format != plan_format
        or marker.get("plan_sha256") != _sha256_path(output / "plan.json")
        or marker.get("jobs_sha256") != _sha256_path(output / "jobs.tsv")
        or marker.get("study_sha256") != _sha256_path(output / "study.json")
        or marker.get("provenance_sha256") != _sha256_path(output / "provenance.json")
    ):
        raise ReductionError(f"invalid or incomplete reduction plan: {output}")
    plan_hash = plan.pop("plan_sha256", None)
    if plan_hash != _hash_json(plan):
        raise ReductionError(f"plan content hash mismatch: {output}")
    plan["plan_sha256"] = plan_hash
    return plan


def _lookup(plan: Mapping[str, object], collection: str, identifier: str) -> dict[str, object]:
    matches = [
        item for item in plan[collection]
        if isinstance(item, dict) and item.get("id") == identifier
    ]
    if len(matches) != 1:
        raise ReductionError(f"plan has no unique {collection} id {identifier}")
    return matches[0]


class _ProcessRegistry:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._processes: set[subprocess.Popen[bytes]] = set()

    def add(self, process: subprocess.Popen[bytes]) -> None:
        with self._lock:
            self._processes.add(process)

    def discard(self, process: subprocess.Popen[bytes]) -> None:
        with self._lock:
            self._processes.discard(process)

    def terminate_all(self, grace: float) -> None:
        with self._lock:
            processes = list(self._processes)
        for process in processes:
            _terminate_process(process, grace)


PROCESS_REGISTRY = _ProcessRegistry()


def _terminate_process(process: subprocess.Popen[bytes], grace: float) -> float:
    started = time.monotonic()
    if process.poll() is not None:
        return 0.0
    try:
        os.killpg(process.pid, signal.SIGTERM)
    except ProcessLookupError:
        return 0.0
    try:
        process.wait(timeout=grace)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.wait()
    return time.monotonic() - started


def _run_command(
    command: Sequence[str], *, cwd: Path, timeout: float, grace: float,
    env: Mapping[str, str] | None = None,
) -> dict[str, object]:
    started = time.monotonic()
    try:
        process = subprocess.Popen(
            list(command), cwd=cwd, env=dict(env) if env is not None else None,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True,
        )
    except OSError as exc:
        return {
            "returncode": None,
            "timed_out": False,
            "wall_sec": time.monotonic() - started,
            "cleanup_wall_sec": 0.0,
            "stdout": b"",
            "stderr": str(exc).encode("utf-8", errors="replace"),
            "error": f"{type(exc).__name__}: {exc}",
        }
    PROCESS_REGISTRY.add(process)
    try:
        try:
            stdout, stderr = process.communicate(timeout=timeout)
            return {
                "returncode": process.returncode,
                "timed_out": False,
                "wall_sec": time.monotonic() - started,
                "cleanup_wall_sec": 0.0,
                "stdout": stdout,
                "stderr": stderr,
                "error": None,
            }
        except subprocess.TimeoutExpired:
            cleanup = _terminate_process(process, grace)
            stdout, stderr = process.communicate()
            return {
                "returncode": process.returncode,
                "timed_out": True,
                "wall_sec": time.monotonic() - started,
                "cleanup_wall_sec": cleanup,
                "stdout": stdout,
                "stderr": stderr,
                "error": None,
            }
    finally:
        PROCESS_REGISTRY.discard(process)


def _observation(result: Mapping[str, object]) -> dict[str, object]:
    stdout = bytes(result["stdout"])
    stderr = bytes(result["stderr"])
    return {
        "returncode": result["returncode"],
        "timed_out": result["timed_out"],
        "wall_sec": result["wall_sec"],
        "stdout_bytes": len(stdout),
        "stderr_bytes": len(stderr),
        "stdout_sha256": _sha256_bytes(stdout),
        "stderr_sha256": _sha256_bytes(stderr),
        "stdout": stdout.decode("utf-8", errors="replace"),
        "stderr": stderr.decode("utf-8", errors="replace"),
        "error": result["error"],
    }


def _match_args(match: Mapping[str, object]) -> list[str]:
    result = []
    if match.get("ignore_stdout"):
        result.append("--ignore-stdout")
    if match.get("ignore_stderr"):
        result.append("--ignore-stderr")
    if match.get("match_stdout") is not None:
        result.extend(("--match-stdout", str(match["match_stdout"])))
    if match.get("match_stderr") is not None:
        result.extend(("--match-stderr", str(match["match_stderr"])))
    return result


def _render_wrapper(
    plan: Mapping[str, object], benchmark: Mapping[str, object], *, journal: Path,
    phase: str, job_id: str, attempt: int,
) -> tuple[list[str], dict[str, str]]:
    wrapper = plan["predicate_wrapper"]
    values = {
        "journal": str(journal),
        "phase": phase,
        "predicate_timeout": f"{float(plan['limits']['predicate_timeout_sec']):g}",
        "run_id": str(plan.get("run_id", plan["study_id"])),
        "job_id": job_id,
        "attempt": str(attempt),
    }
    command: list[str] = []
    for token in wrapper["command"]:
        if token == "{command}":
            command.extend(str(item) for item in benchmark["predicate"]["command"])
        elif token == "{match_args}":
            command.extend(_match_args(benchmark["predicate"]["match"]))
        else:
            try:
                command.append(str(token).format_map(values))
            except KeyError as exc:
                raise ReductionError(f"unknown predicate wrapper placeholder: {exc.args[0]}") from exc
    env = dict(os.environ)
    env.update(wrapper.get("env", {}))
    env.update({
        "SMTBATCH_RUN_ID": str(plan.get("run_id", plan["study_id"])),
        "SMTBATCH_JOB_ID": job_id,
        "SMTBATCH_ATTEMPT": str(attempt),
        "SMTBATCH_PHASE": phase,
        "SMTBATCH_INPUT_SHA256": str(benchmark["input_sha256"]),
    })
    return command, env


def _render_tool_command(
    plan: Mapping[str, object], benchmark: Mapping[str, object],
    reducer: Mapping[str, object], job: Mapping[str, object], attempt_dir: Path,
) -> tuple[list[str], dict[str, str]]:
    journal = attempt_dir / "predicate.jsonl"
    predicate, wrapper_env = _render_wrapper(
        plan, benchmark, journal=journal, phase="reducer",
        job_id=str(job["job_id"]), attempt=int(job["attempt"]),
    )
    limits = plan["limits"]
    values = {
        "input": str(benchmark["input"]),
        "output": str(attempt_dir / "output.smt2"),
        "workdir": str(attempt_dir),
        "predicate_timeout": f"{float(limits['predicate_timeout_sec']):g}",
    }
    command: list[str] = []
    for token in reducer["command"]:
        if token == "{predicate}":
            command.extend(predicate)
        else:
            try:
                command.append(str(token).format_map(values))
            except KeyError as exc:
                raise ReductionError(f"unknown tool command placeholder: {exc.args[0]}") from exc
    env = wrapper_env
    env.update(reducer.get("env", {}))
    env.update({
        "PYTHONUNBUFFERED": "1",
        "SMTBATCH_REDUCER_ID": str(reducer["id"]),
        "SMTBATCH_REPEAT": str(job["repeat"]),
    })
    return command, env


def _run_predicate(
    plan: Mapping[str, object], benchmark: Mapping[str, object], candidate: Path,
    *, journal: Path, phase: str, job_id: str, attempt: int,
) -> dict[str, object]:
    command, env = _render_wrapper(
        plan, benchmark, journal=journal, phase=phase, job_id=job_id, attempt=attempt
    )
    result = _run_command(
        [*command, str(candidate)],
        cwd=Path(str(plan["root"])),
        timeout=float(plan["limits"]["predicate_timeout_sec"]) + 5,
        grace=float(plan["limits"]["termination_grace_sec"]),
        env=env,
    )
    return _observation(result)


def _stream_matches(
    baseline: Mapping[str, object], candidate: Mapping[str, object], *, stream: str,
    match: Mapping[str, object],
) -> bool:
    if match.get(f"ignore_{stream}"):
        return True
    expected = match.get(f"match_{stream}")
    if expected is not None:
        return str(expected) in str(candidate.get(stream, ""))
    return candidate.get(f"{stream}_sha256") == baseline.get(f"{stream}_sha256")


def _matches_baseline(
    baseline: Mapping[str, object], candidate: Mapping[str, object], match: Mapping[str, object]
) -> bool:
    if baseline.get("error") is not None or candidate.get("error") is not None:
        return False
    if bool(baseline.get("timed_out")) != bool(candidate.get("timed_out")):
        return False
    if baseline.get("returncode") != candidate.get("returncode"):
        return False
    return _stream_matches(baseline, candidate, stream="stdout", match=match) and _stream_matches(
        baseline, candidate, stream="stderr", match=match
    )


def _read_jsonl(path: Path, *, allow_partial: bool) -> tuple[list[dict[str, object]], bool, str]:
    if not path.is_file():
        return [], False, "missing"
    try:
        data = path.read_bytes()
    except OSError as exc:
        return [], False, f"{type(exc).__name__}: {exc}"
    rows = []
    truncated = False
    lines = data.splitlines(keepends=True)
    for index, raw in enumerate(lines):
        if not raw.strip():
            continue
        complete = raw.endswith((b"\n", b"\r"))
        try:
            value = json.loads(raw.decode("utf-8"))
        except (UnicodeError, json.JSONDecodeError) as exc:
            if allow_partial and index == len(lines) - 1 and not complete:
                truncated = True
                break
            return rows, truncated, f"line {index + 1}: {exc}"
        if not isinstance(value, dict):
            return rows, truncated, f"line {index + 1}: event is not an object"
        rows.append(value)
    return rows, truncated, ""


def _journal_events(path: Path, *, allow_partial: bool) -> dict[str, object]:
    rows, truncated, error = _read_jsonl(path, allow_partial=allow_partial)
    starts: list[dict[str, object]] = []
    finishes: dict[str, dict[str, object]] = {}
    seen_starts: set[str] = set()
    previous_seq = 0
    for row in rows:
        event = row.get("event")
        call_id = row.get("call_id")
        if event not in {"start", "finish"} or not isinstance(call_id, str):
            error = error or "journal contains an invalid event"
            continue
        sequence = row.get("call_seq")
        if isinstance(sequence, int) and not isinstance(sequence, bool):
            if event == "start" and sequence <= previous_seq:
                error = error or "call_seq is not strictly increasing"
            if event == "start":
                previous_seq = sequence
        if event == "start":
            if call_id in seen_starts:
                error = error or f"duplicate start {call_id}"
            starts.append(row)
            seen_starts.add(call_id)
        else:
            if call_id in finishes:
                error = error or f"duplicate finish {call_id}"
            finishes[call_id] = row
    incomplete = [row["call_id"] for row in starts if row["call_id"] not in finishes]
    return {
        "rows": rows,
        "starts": starts,
        "finishes": finishes,
        "truncated_tail": truncated,
        "error": error,
        "incomplete": incomplete,
    }


def _quality(value: object) -> tuple[int, int, int] | None:
    if not isinstance(value, dict):
        return None
    names = ("expression_count", "node_count", "byte_count")
    values = tuple(value.get(name) for name in names)
    if not all(isinstance(item, int) and not isinstance(item, bool) and item >= 0 for item in values):
        return None
    return values  # type: ignore[return-value]


def _candidate_quality(start: Mapping[str, object]) -> tuple[int, int, int] | None:
    candidate = start.get("candidate")
    return _quality(candidate.get("quality")) if isinstance(candidate, dict) else None


def _candidate_hash(start: Mapping[str, object]) -> str:
    candidate = start.get("candidate")
    if not isinstance(candidate, dict):
        return ""
    return str(candidate.get("canonical_sha256") or candidate.get("sha256") or "")


def trajectory_for_attempt(
    attempt_dir: Path, benchmark: Mapping[str, object], reducer: Mapping[str, object],
    *, allow_partial: bool,
) -> dict[str, object]:
    journal = _journal_events(attempt_dir / "predicate.jsonl", allow_partial=allow_partial)
    starts = list(journal["starts"])
    finishes = journal["finishes"]
    reducer_starts = [row for row in starts if row.get("phase") == "reducer"]
    golden = reducer_starts[0] if reducer_starts else None
    initial = _candidate_quality(golden) if golden else None
    if initial is None:
        initial = (0, 0, int(benchmark.get("input_bytes", 0)))
    golden_finish = finishes.get(golden["call_id"]) if golden else None
    candidates = reducer_starts[1:]
    phase_boundaries = []
    exact = False
    best = initial
    points = [{
        "call_index": 0, "elapsed_sec": 0.0,
        "expression_count": initial[0], "node_count": initial[1], "byte_count": initial[2],
        "accepted": True, "preserving": True, "candidate_sha256": _candidate_hash(golden or {}),
        "phase": "initial", "mutator": None,
    }]
    accepted_count = 0
    preserving_count = 0
    start_ns = int((golden or {}).get("monotonic_ns", 0) or 0)
    match = benchmark["predicate"]["match"]
    warnings = []
    for call_index, start in enumerate(candidates, 1):
        finish = finishes.get(start["call_id"])
        preserving = bool(
            golden_finish is not None and finish is not None
            and _matches_journal_finish(golden_finish, finish, match)
        )
        if preserving:
            preserving_count += 1
        accepted = preserving
        candidate_quality = _candidate_quality(start)
        candidate = start.get("candidate")
        if candidate_quality is None:
            detail = candidate.get("quality_error") if isinstance(candidate, dict) else None
            warnings.append(
                "candidate size vector is unavailable"
                + (f": {detail}" if detail else "")
            )
        if accepted and candidate_quality is not None:
            best = candidate_quality
            accepted_count += 1
        finish_ns = int((finish or start).get("monotonic_ns", 0) or 0)
        elapsed = max(0.0, (finish_ns - start_ns) / 1_000_000_000) if start_ns and finish_ns else 0.0
        points.append({
            "call_index": call_index,
            "elapsed_sec": elapsed,
            "expression_count": best[0], "node_count": best[1], "byte_count": best[2],
            "candidate_expression_count": candidate_quality[0] if candidate_quality else None,
            "candidate_node_count": candidate_quality[1] if candidate_quality else None,
            "candidate_byte_count": candidate_quality[2] if candidate_quality else None,
            "accepted": accepted, "preserving": preserving,
            "candidate_sha256": _candidate_hash(start),
            "predicate_call_id": start["call_id"],
            "phase": None,
            "pass": None,
            "mutator": None,
            "size_delta": None,
        })
    if golden is None:
        warnings.append("reducer golden predicate call is missing")
    elif _candidate_quality(golden) is None:
        warnings.append("reducer golden size vector is unavailable")
    if journal["error"]:
        warnings.append(str(journal["error"]))
    if journal["incomplete"] and not allow_partial:
        warnings.append(f"{len(journal['incomplete'])} predicate calls lack finish events")
    if journal["truncated_tail"] and not allow_partial:
        warnings.append("evidence has a truncated tail")
    last_accept = max((point["call_index"] for point in points if point["accepted"]), default=0)
    return {
        "provisional": allow_partial,
        "exact_acceptance": exact,
        "initial": {"expression_count": initial[0], "node_count": initial[1], "byte_count": initial[2]},
        "final": {"expression_count": best[0], "node_count": best[1], "byte_count": best[2]},
        "candidate_calls": len(candidates),
        "preserving_calls": preserving_count,
        "accepted_moves": accepted_count,
        "last_accept_call": last_accept,
        "tail_calls": max(0, len(candidates) - last_accept),
        "points": points,
        "phase_boundaries": phase_boundaries,
        "warnings": warnings,
        "evidence_ok": not warnings,
    }


def _matches_journal_finish(
    golden: Mapping[str, object], candidate: Mapping[str, object], match: Mapping[str, object]
) -> bool:
    if bool(golden.get("timed_out")) != bool(candidate.get("timed_out")):
        return False
    if candidate.get("returncode") != golden.get("returncode"):
        return False
    for stream in ("stdout", "stderr"):
        if match.get(f"ignore_{stream}"):
            continue
        if match.get(f"match_{stream}") is not None:
            if candidate.get(f"{stream}_match") is not True:
                return False
        elif candidate.get(f"{stream}_sha256") != golden.get(f"{stream}_sha256"):
            return False
    return True


def _artifact_records(directory: Path) -> list[dict[str, object]]:
    records = []
    ignored = {"artifacts.json", "attempt.complete.json"}
    for path in sorted(directory.rglob("*"), key=lambda item: item.relative_to(directory).as_posix()):
        if path.is_symlink():
            raise ReductionError(f"attempt artifact must not be a symlink: {path}")
        if not path.is_file() or path.name in ignored:
            continue
        records.append({
            "path": path.relative_to(directory).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": _sha256_path(path),
        })
    return records


def _marker_valid(job_dir: Path) -> bool:
    marker_path = job_dir / "job.complete.json"
    result_path = job_dir / "result.json"
    if not marker_path.is_file() or not result_path.is_file():
        return False
    try:
        marker = _mapping(_read_json(marker_path, "job marker"), "job marker")
        return (
            marker.get("schema_version") in SUPPORTED_FORMATS
            and marker.get("result_sha256") == _sha256_path(result_path)
        )
    except (OSError, ReductionError):
        return False


def _next_attempt(job_dir: Path) -> tuple[int, Path]:
    attempts = job_dir / "attempts"
    attempts.mkdir(parents=True, exist_ok=True)
    existing = [int(path.name) for path in attempts.iterdir() if path.is_dir() and path.name.isdigit()]
    number = max(existing, default=0) + 1
    path = attempts / f"{number:04d}"
    path.mkdir()
    return number, path


def _seal_job(job_dir: Path, attempt_dir: Path, result: dict[str, object]) -> None:
    _write_json(attempt_dir / "result.json", result)
    _write_json(attempt_dir / "artifacts.json", {
        "schema_version": SCHEMA_VERSION,
        "files": _artifact_records(attempt_dir),
    })
    _write_json(attempt_dir / "attempt.complete.json", {
        "schema_version": SCHEMA_VERSION,
        "result_sha256": _sha256_path(attempt_dir / "result.json"),
        "artifacts_sha256": _sha256_path(attempt_dir / "artifacts.json"),
    })
    _write_json(job_dir / "result.json", result)
    _write_json(job_dir / "job.complete.json", {
        "schema_version": SCHEMA_VERSION,
        "result_sha256": _sha256_path(job_dir / "result.json"),
    })


def _control_path(output: Path) -> Path:
    return output / "control.json"


def _control_mode(output: Path) -> str:
    path = _control_path(output)
    if not path.is_file():
        return ""
    try:
        value = _mapping(_read_json(path, "run control"), "run control")
    except ReductionError:
        return "immediate"
    mode = value.get("mode")
    return str(mode) if mode in {"graceful", "immediate"} else "immediate"


def request_stop(output: Path, mode: str) -> dict[str, object]:
    if mode not in {"graceful", "immediate"}:
        raise ReductionError("stop mode must be graceful or immediate")
    output = output.expanduser().resolve()
    plan = load_plan(output)
    version, _ = _schema_identity(plan, "plan")
    if version != SCHEMA_VERSION:
        raise ReductionError("reduction-v2 plans are read-only and cannot be stopped")
    current = _control_mode(output)
    effective = "immediate" if "immediate" in {current, mode} else "graceful"
    request = {
        "schema_version": SCHEMA_VERSION,
        "mode": effective,
        "requested_at": _utc_now(),
        "requested_by_pid": os.getpid(),
    }
    _write_json(_control_path(output), request)
    _append_jsonl(output / "control_history.jsonl", {"event": "stop_requested", **request})
    return request


def _clear_control_for_resume(output: Path) -> None:
    path = _control_path(output)
    if path.is_file():
        try:
            previous = _read_json(path, "run control")
        except ReductionError as exc:
            previous = {"error": str(exc)}
        _append_jsonl(output / "control_history.jsonl", {
            "event": "control_cleared_for_resume", "at": _utc_now(), "value": previous,
        })
        path.unlink()


def _evidence_health(
    attempt_dir: Path, benchmark: Mapping[str, object], reducer: Mapping[str, object],
    *, allow_partial: bool,
) -> dict[str, object]:
    trajectory = trajectory_for_attempt(
        attempt_dir, benchmark, reducer, allow_partial=allow_partial
    )
    warnings = list(trajectory["warnings"])
    return {
        "ok": not warnings,
        "warnings": list(dict.fromkeys(warnings)),
        "predicate_calls": trajectory["candidate_calls"],
        "accepted_moves": trajectory["accepted_moves"],
        "trajectory": trajectory,
    }


def _partial_abort(attempt_dir: Path, reason: str) -> None:
    try:
        artifacts = _artifact_records(attempt_dir)
    except ReductionError:
        artifacts = []
    _write_json(attempt_dir / "partial.json", {
        "schema_version": SCHEMA_VERSION,
        "status": "aborted",
        "reason": reason,
        "at": _utc_now(),
        "artifacts": artifacts,
    })


def _execute_job(output: Path, plan: Mapping[str, object], job: Mapping[str, object]) -> dict[str, object]:
    job_dir = output / "jobs" / str(job["job_id"])
    job_dir.mkdir(parents=True, exist_ok=True)
    if _marker_valid(job_dir):
        return _mapping(_read_json(job_dir / "result.json", "job result"), "job result")
    claim = job_dir / ".claim"
    if claim.is_dir():
        try:
            claim.rmdir()
        except OSError as exc:
            raise ReductionError(f"stale job claim is not empty: {claim}") from exc
    try:
        claim.mkdir()
    except FileExistsError as exc:
        raise ReductionError(f"job is already claimed: {job['job_id']}") from exc
    try:
        attempt, attempt_dir = _next_attempt(job_dir)
        enriched_job = {**job, "attempt": attempt}
        benchmark = _lookup(plan, "benchmarks", str(job["benchmark_id"]))
        reducer = _lookup(plan, "reducers", str(job["reducer_id"]))
        limits = plan["limits"]
        root = Path(str(plan["root"]))
        input_path = Path(str(benchmark["input"]))
        journal = attempt_dir / "predicate.jsonl"
        if not input_path.is_file() or _sha256_path(input_path) != benchmark["input_sha256"]:
            raise ReductionError(f"benchmark input drift: {input_path}")
        _write_json(attempt_dir / "job.json", enriched_job)

        preflight = []
        for _ in range(int(limits["preflight_repeats"])):
            preflight.append(_run_predicate(
                plan, benchmark, input_path, journal=journal, phase="preflight",
                job_id=str(job["job_id"]), attempt=attempt,
            ))
            if _control_mode(output) == "immediate":
                _partial_abort(attempt_dir, "immediate stop during preflight")
                raise ImmediateAbort("immediate stop during preflight")
        _write_json(attempt_dir / "preflight.json", {"observations": preflight})
        stable_preflight = bool(preflight) and all(
            _matches_baseline(preflight[0], item, benchmark["predicate"]["match"])
            for item in preflight
        )
        if not stable_preflight:
            result = {
                **enriched_job,
                "schema_version": SCHEMA_VERSION, "format": FORMAT,
                "status": "invalid", "status_detail": "preflight_failed",
                "verified": False,
                "evidence_ok": False, "evidence_warnings": ["preflight is unstable"],
                "input_quality": None, "output_quality": None,
                "input_bytes": benchmark["input_bytes"], "output_bytes": None,
                "size_ratio": None, "trial_wall_sec": 0.0, "cleanup_wall_sec": 0.0,
                "predicate_calls": 0, "accepted_moves": 0, "output": "",
                "attempt_dir": str(attempt_dir), "finished_at": _utc_now(),
            }
            _seal_job(job_dir, attempt_dir, result)
            return result

        # A reducer is allowed to reach a fixed point without writing its
        # output path.  Materialize the incumbent before launch so that
        # "no reduction" is represented by a verified original candidate,
        # while a crash/timeout still receives its distinct reducer status.
        shutil.copy2(input_path, attempt_dir / "output.smt2")
        command, env = _render_tool_command(plan, benchmark, reducer, enriched_job, attempt_dir)
        _write_json(attempt_dir / "command.json", {
            "argv": command, "cwd": str(root),
            "env": {key: env[key] for key in sorted(set(reducer.get("env", {})) | {
                "SMTBATCH_RUN_ID", "SMTBATCH_JOB_ID", "SMTBATCH_ATTEMPT",
                "SMTBATCH_REDUCER_ID", "SMTBATCH_REPEAT",
            }) if key in env},
        })
        reducer_result = _run_command(
            command, cwd=root, timeout=float(limits["trial_wall_sec"]),
            grace=float(limits["termination_grace_sec"]), env=env,
        )
        reducer_stdout = bytes(reducer_result.pop("stdout"))
        reducer_stderr = bytes(reducer_result.pop("stderr"))
        (attempt_dir / "reducer.stdout").write_bytes(reducer_stdout)
        (attempt_dir / "reducer.stderr").write_bytes(reducer_stderr)
        _write_json(attempt_dir / "reducer.json", reducer_result)
        if _control_mode(output) == "immediate":
            _partial_abort(attempt_dir, "immediate stop during reducer")
            raise ImmediateAbort("immediate stop during reducer")

        output_path = attempt_dir / "output.smt2"
        output_valid = output_path.is_file() and not output_path.is_symlink()
        verification = []
        verified = False
        if output_valid:
            for _ in range(int(limits["verification_repeats"])):
                verification.append(_run_predicate(
                    plan, benchmark, output_path, journal=journal, phase="final-replay",
                    job_id=str(job["job_id"]), attempt=attempt,
                ))
            verified = bool(verification) and all(
                _matches_baseline(preflight[0], item, benchmark["predicate"]["match"])
                for item in verification
            )
        _write_json(attempt_dir / "verification.json", {
            "output_valid": output_valid, "verified": verified, "observations": verification,
        })
        if verified:
            shutil.copy2(output_path, attempt_dir / "output.verified.smt2")

        if reducer_result["error"] is not None:
            status_value, status_detail = "invalid", "infrastructure_error"
        elif not output_valid:
            status_value, status_detail = "invalid", "invalid_output"
        elif not verified:
            status_value, status_detail = "invalid", "verification_failed"
        elif reducer_result["timed_out"]:
            status_value, status_detail = "truncated", "timeout"
        elif reducer_result["returncode"] != 0:
            status_value, status_detail = "truncated", "reducer_error"
        else:
            status_value, status_detail = "completed", ""
        health = _evidence_health(attempt_dir, benchmark, reducer, allow_partial=False)
        trajectory = health["trajectory"]
        initial_quality = trajectory["initial"]
        output_quality = trajectory["final"]
        output_bytes = output_path.stat().st_size if output_valid else None
        result = {
            **enriched_job,
            "schema_version": SCHEMA_VERSION, "format": FORMAT,
            "status": status_value, "status_detail": status_detail,
            "verified": verified,
            "evidence_ok": health["ok"], "evidence_warnings": health["warnings"],
            "input_quality": initial_quality, "output_quality": output_quality,
            "input_bytes": benchmark["input_bytes"], "input_sha256": benchmark["input_sha256"],
            "output_bytes": output_bytes,
            "output_sha256": _sha256_path(output_path) if output_valid else None,
            "size_ratio": output_bytes / benchmark["input_bytes"] if output_bytes is not None and benchmark["input_bytes"] else None,
            "output_changed": output_valid and _sha256_path(output_path) != benchmark["input_sha256"],
            "trial_wall_sec": reducer_result["wall_sec"],
            "cleanup_wall_sec": reducer_result["cleanup_wall_sec"],
            "returncode": reducer_result["returncode"], "timed_out": reducer_result["timed_out"],
            "predicate_calls": health["predicate_calls"], "accepted_moves": health["accepted_moves"],
            "output": str(attempt_dir / "output.verified.smt2") if verified else "",
            "attempt_dir": str(attempt_dir), "finished_at": _utc_now(),
        }
        _seal_job(job_dir, attempt_dir, result)
        return result
    finally:
        try:
            claim.rmdir()
        except OSError:
            pass


def completed_results(output: Path, plan: Mapping[str, object]) -> list[dict[str, object]]:
    results = []
    for job in plan["jobs"]:
        job_dir = output / "jobs" / str(job["job_id"])
        if _marker_valid(job_dir):
            results.append(_mapping(_read_json(job_dir / "result.json", "job result"), "job result"))
    return results


def _quality_columns(value: object, prefix: str) -> dict[str, object]:
    quality = value if isinstance(value, dict) else {}
    return {
        f"{prefix}_expressions": quality.get("expression_count"),
        f"{prefix}_nodes": quality.get("node_count"),
        f"{prefix}_bytes": quality.get("byte_count"),
    }


def write_results_index(output: Path, plan: Mapping[str, object]) -> list[dict[str, object]]:
    results = completed_results(output, plan)
    temporary = output / ".results.tsv.tmp"
    try:
        with temporary.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=RESULT_FIELDS, delimiter="\t")
            writer.writeheader()
            for result in sorted(results, key=lambda item: int(item["order"])):
                row = {
                    "job_id": result["job_id"], "benchmark": result["benchmark_id"],
                    "reducer": result["reducer_id"], "repeat": result["repeat"],
                    "wave": result["wave"], "status": result["status"],
                    "status_detail": result.get("status_detail", ""),
                    "verified": str(bool(result["verified"])).lower(),
                    "evidence_ok": str(bool(result["evidence_ok"])).lower(),
                    **_quality_columns(result.get("input_quality"), "input"),
                    **_quality_columns(result.get("output_quality"), "output"),
                    "size_ratio": result.get("size_ratio") if result.get("size_ratio") is not None else "",
                    "trial_wall_sec": result.get("trial_wall_sec", 0),
                    "cleanup_wall_sec": result.get("cleanup_wall_sec", 0),
                    "predicate_calls": result.get("predicate_calls", 0),
                    "accepted_moves": result.get("accepted_moves", 0),
                    "attempt": result["attempt"], "output": result.get("output", ""),
                }
                writer.writerow(row)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.replace(output / "results.tsv")
    finally:
        temporary.unlink(missing_ok=True)
    return results


def _progress(
    output: Path, plan: Mapping[str, object], status_value: str,
    running_jobs: Sequence[Mapping[str, object]] = (),
) -> None:
    results = completed_results(output, plan)
    statuses = Counter(str(item["status"]) for item in results)
    by_reducer: dict[str, dict[str, object]] = {}
    for reducer in plan["reducers"]:
        selected = [item for item in results if item["reducer_id"] == reducer["id"]]
        by_reducer[str(reducer["id"])] = {
            "label": reducer["label"],
            "completed": len(selected),
            "verified": sum(bool(item.get("verified")) for item in selected),
            "evidence_ok": sum(bool(item.get("evidence_ok")) for item in selected),
            "predicate_calls": sum(int(item.get("predicate_calls", 0)) for item in selected),
            "statuses": dict(Counter(str(item["status"]) for item in selected)),
        }
    running = [
        {
            "job_id": item["job_id"], "benchmark": item["benchmark_id"],
            "reducer": item["reducer_id"], "repeat": item["repeat"], "wave": item["wave"],
        }
        for item in running_jobs
    ]
    _write_json(output / "progress.json", {
        "schema_version": SCHEMA_VERSION, "format": FORMAT,
        "study_id": plan["study_id"], "status": status_value, "updated_at": _utc_now(),
        "total_jobs": len(plan["jobs"]), "completed_jobs": len(results),
        "running_jobs": len(running), "pending_jobs": len(plan["jobs"]) - len(results) - len(running),
        "statuses": dict(statuses), "by_reducer": by_reducer, "active": running,
    })


def _append_resume(output: Path, event: str, **values: object) -> None:
    _append_jsonl(output / "resume_history.jsonl", {
        "schema_version": SCHEMA_VERSION, "event": event, "at": _utc_now(),
        "pid": os.getpid(), **values,
    })


def _run_locked(output: Path, plan: Mapping[str, object]) -> list[dict[str, object]]:
    # The prepared plan remains immutable on disk, while subprocess evidence
    # still needs the concrete run directory identity.
    plan = dict(plan)
    plan["run_id"] = output.name
    _clear_control_for_resume(output)
    workers = int(plan["execution"]["outer_jobs"])
    _append_resume(
        output, "run_started", outer_jobs=workers,
        completed_before=len(completed_results(output, plan)),
    )
    _progress(output, plan, "running")
    stopped = ""
    active: dict[concurrent.futures.Future[dict[str, object]], Mapping[str, object]] = {}

    def signal_stop(signum: int, _frame: object) -> None:
        request_stop(output, "graceful" if signum == signal.SIGINT else "immediate")

    previous_int = signal.getsignal(signal.SIGINT)
    previous_term = signal.getsignal(signal.SIGTERM)
    signal.signal(signal.SIGINT, signal_stop)
    signal.signal(signal.SIGTERM, signal_stop)
    try:
        waves = sorted({int(job["wave"]) for job in plan["jobs"]})
        for wave in waves:
            pending = [
                job for job in plan["jobs"]
                if int(job["wave"]) == wave
                and not _marker_valid(output / "jobs" / str(job["job_id"]))
            ]
            if not pending:
                continue
            pending_iter = iter(pending)
            exhausted = False
            with concurrent.futures.ThreadPoolExecutor(
                max_workers=workers, thread_name_prefix="reduction-job"
            ) as executor:
                while active or not exhausted:
                    mode = _control_mode(output)
                    if mode:
                        stopped = mode
                    if mode == "immediate":
                        _progress(output, plan, "aborting", list(active.values()))
                        PROCESS_REGISTRY.terminate_all(float(plan["limits"]["termination_grace_sec"]))
                    while not stopped and len(active) < workers and not exhausted:
                        try:
                            job = next(pending_iter)
                        except StopIteration:
                            exhausted = True
                            break
                        future = executor.submit(_execute_job, output, plan, job)
                        active[future] = job
                    if not active:
                        break
                    _progress(
                        output, plan, "stopping" if stopped == "graceful" else "running",
                        list(active.values()),
                    )
                    done, _ = concurrent.futures.wait(
                        active, timeout=0.2, return_when=concurrent.futures.FIRST_COMPLETED
                    )
                    for future in done:
                        active.pop(future)
                        try:
                            future.result()
                        except ImmediateAbort:
                            stopped = "immediate"
                        write_results_index(output, plan)
                    if stopped and not active:
                        break
            if stopped:
                break
        results = write_results_index(output, plan)
        if stopped:
            _progress(output, plan, "interrupted")
            _append_resume(
                output, "run_finished", status="interrupted", stop_mode=stopped,
                completed_jobs=len(results), remaining_jobs=len(plan["jobs"]) - len(results),
            )
        else:
            _progress(output, plan, "complete")
            _append_resume(
                output, "run_finished", status="complete", completed_jobs=len(results), remaining_jobs=0,
            )
        return results
    except Exception as exc:
        PROCESS_REGISTRY.terminate_all(float(plan["limits"]["termination_grace_sec"]))
        write_results_index(output, plan)
        _progress(output, plan, "failed")
        _append_resume(
            output, "run_finished", status="failed",
            error=f"{type(exc).__name__}: {exc}",
            completed_jobs=len(completed_results(output, plan)),
        )
        raise
    finally:
        signal.signal(signal.SIGINT, previous_int)
        signal.signal(signal.SIGTERM, previous_term)


def run(output: Path) -> list[dict[str, object]]:
    output = output.expanduser().resolve()
    plan = load_plan(output)
    version, _ = _schema_identity(plan, "plan")
    if version != SCHEMA_VERSION:
        raise ReductionError(
            "reduction-v2 plans are read-only; prepare a reduction-v3 plan before execution"
        )
    lock = _RunLock(output)
    try:
        lock.acquire()
        try:
            _validate_live_provenance(plan)
        except ReductionError as exc:
            _append_jsonl(output / "resume_history.jsonl", {
                "schema_version": SCHEMA_VERSION,
                "event": "run_rejected",
                "at": _utc_now(),
                "pid": os.getpid(),
                "reason": str(exc),
            })
            raise
        return _run_locked(output, plan)
    except ValueError as exc:
        raise ReductionError(str(exc)) from exc
    finally:
        lock.release()


def status(output: Path) -> dict[str, object]:
    output = output.expanduser().resolve()
    plan = load_plan(output)
    results = completed_results(output, plan)
    progress = {}
    try:
        progress = _mapping(_read_json(output / "progress.json", "progress"), "progress")
    except ReductionError:
        pass
    return {
        "study_id": plan["study_id"], "format": plan["format"],
        "benchmarks": len(plan["benchmarks"]), "reducers": len(plan["reducers"]),
        "repeats": plan["repeats"], "outer_jobs": plan["execution"]["outer_jobs"],
        "selection": plan.get("selection", {}), "limits": plan["limits"],
        "total_jobs": len(plan["jobs"]), "completed_jobs": len(results),
        "pending_jobs": len(plan["jobs"]) - len(results),
        "statuses": dict(Counter(str(item["status"]) for item in results)),
        "status": progress.get("status", "prepared"),
    }


def _latest_attempt(job_dir: Path) -> Path | None:
    attempts = job_dir / "attempts"
    if not attempts.is_dir():
        return None
    values = sorted((path for path in attempts.iterdir() if path.is_dir() and path.name.isdigit()))
    return values[-1] if values else None


def _attempt_log(attempt_dir: Path, name: str) -> dict[str, object]:
    path = attempt_dir / name
    try:
        data = path.read_bytes()
    except OSError as exc:
        return {"text": "", "truncated": False, "error": str(exc)}
    truncated = len(data) > MAX_LOG_BYTES
    if truncated:
        marker = b"\n\n... log truncated by SMTBatch ...\n\n"
        half = max(0, (MAX_LOG_BYTES - len(marker)) // 2)
        data = (
            data[:half]
            + marker
            + data[-half:]
        )
    return {
        "text": data.decode("utf-8", errors="replace"),
        "truncated": truncated,
        "error": None,
    }


def trajectory_for_case(output: Path, case_id: str) -> dict[str, object]:
    output = output.expanduser().resolve()
    plan = load_plan(output)
    benchmark = _lookup(plan, "benchmarks", case_id)
    trials = []
    for job in plan["jobs"]:
        if job["benchmark_id"] != case_id:
            continue
        job_dir = output / "jobs" / str(job["job_id"])
        sealed = _marker_valid(job_dir)
        result = (
            _mapping(_read_json(job_dir / "result.json", "job result"), "job result")
            if sealed else {}
        )
        attempt_dir = Path(str(result.get("attempt_dir", ""))) if result else _latest_attempt(job_dir)
        reducer = _lookup(plan, "reducers", str(job["reducer_id"]))
        if attempt_dir is None or not attempt_dir.is_dir():
            trajectory = {
                "provisional": True, "exact_acceptance": False,
                "initial": None, "final": None, "candidate_calls": 0,
                "preserving_calls": 0, "accepted_moves": 0,
                "last_accept_call": 0, "tail_calls": 0,
                "points": [], "phase_boundaries": [], "warnings": ["trial has not started"],
                "evidence_ok": False,
            }
        else:
            trajectory = trajectory_for_attempt(
                attempt_dir, benchmark, reducer, allow_partial=not sealed
            )
        logs = {
            "stdout": _attempt_log(attempt_dir, "reducer.stdout")
            if attempt_dir and attempt_dir.is_dir()
            else {"text": "", "truncated": False, "error": None},
            "stderr": _attempt_log(attempt_dir, "reducer.stderr")
            if attempt_dir and attempt_dir.is_dir()
            else {"text": "", "truncated": False, "error": None},
        }
        trials.append({
            "job_id": job["job_id"], "reducer_id": reducer["id"],
            "reducer_label": reducer["label"], "repeat": job["repeat"],
            "wave": job["wave"], "status": result.get("status", "running" if attempt_dir else "pending"),
            "verified": result.get("verified"), "evidence_ok": result.get("evidence_ok"),
            "attempt": result.get("attempt", int(attempt_dir.name) if attempt_dir and attempt_dir.name.isdigit() else None),
            "logs": logs,
            "trajectory": trajectory,
        })
    return {
        "schema_version": plan["schema_version"], "format": plan["format"],
        "study_id": plan["study_id"],
        "case": {
            "id": benchmark["id"], "family": benchmark["family"],
            "theory": benchmark["theory"], "predicate_mode": benchmark["predicate_mode"],
            "solver": benchmark["solver"], "input": benchmark["input"],
        },
        "trials": trials,
    }


def case_rows(output: Path) -> list[dict[str, object]]:
    output = output.expanduser().resolve()
    plan = load_plan(output)
    results = {str(item["job_id"]): item for item in completed_results(output, plan)}
    rows = []
    for benchmark in plan["benchmarks"]:
        jobs = [job for job in plan["jobs"] if job["benchmark_id"] == benchmark["id"]]
        selected = [results[str(job["job_id"])] for job in jobs if str(job["job_id"]) in results]
        by_reducer = {}
        for reducer in plan["reducers"]:
            reducer_results = [item for item in selected if item["reducer_id"] == reducer["id"]]
            by_reducer[str(reducer["id"])] = {
                "label": reducer["label"], "planned": int(plan["repeats"]),
                "completed": len(reducer_results),
                "verified": sum(bool(item.get("verified")) for item in reducer_results),
                "evidence_ok": sum(bool(item.get("evidence_ok")) for item in reducer_results),
                "statuses": dict(Counter(str(item["status"]) for item in reducer_results)),
                "predicate_calls": sum(int(item.get("predicate_calls", 0)) for item in reducer_results),
                "accepted_moves": sum(int(item.get("accepted_moves", 0)) for item in reducer_results),
                "final_quality": [item.get("output_quality") for item in reducer_results],
            }
        rows.append({
            "case_id": benchmark["id"], "family": benchmark["family"],
            "theory": benchmark["theory"], "predicate_mode": benchmark["predicate_mode"],
            "solver": benchmark["solver"], "input": benchmark["input"],
            "planned": len(jobs), "completed": len(selected),
            "verified": sum(bool(item.get("verified")) for item in selected),
            "evidence_ok": sum(bool(item.get("evidence_ok")) for item in selected),
            "statuses": dict(Counter(str(item["status"]) for item in selected)),
            "by_reducer": by_reducer,
        })
    return rows


def _median(values: Iterable[object]) -> float | None:
    numbers = [
        float(value) for value in values
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))
    ]
    return statistics.median(numbers) if numbers else None


def _quality_tuple(value: object) -> tuple[int, int, int] | None:
    return _quality(value)


def _case_reducer_quality(
    results: Sequence[Mapping[str, object]], case_id: str, reducer_id: str, repeats: int
) -> tuple[int, int, int] | None:
    values = sorted(
        value for value in (
            _quality_tuple(item.get("output_quality")) for item in results
            if item.get("benchmark_id") == case_id and item.get("reducer_id") == reducer_id
            and item.get("verified") is True and item.get("evidence_ok") is True
        ) if value is not None
    )
    if len(values) != repeats:
        return None
    return values[(len(values) - 1) // 2]


def comparison_rows(
    plan: Mapping[str, object], results: Sequence[Mapping[str, object]]
) -> list[dict[str, object]]:
    comparisons = []
    for left, right in plan["comparisons"]:
        wins = ties = losses = paired = 0
        for benchmark in plan["benchmarks"]:
            left_value = _case_reducer_quality(
                results, str(benchmark["id"]), str(left), int(plan["repeats"])
            )
            right_value = _case_reducer_quality(
                results, str(benchmark["id"]), str(right), int(plan["repeats"])
            )
            if left_value is None or right_value is None:
                continue
            paired += 1
            if left_value < right_value:
                wins += 1
            elif left_value > right_value:
                losses += 1
            else:
                ties += 1
        comparisons.append({
            "left": left, "right": right, "paired_cases": paired,
            "wins": wins, "ties": ties, "losses": losses,
        })
    return comparisons


def build_report(output: Path) -> tuple[dict[str, object], list[dict[str, object]], list[dict[str, object]]]:
    output = output.expanduser().resolve()
    plan = load_plan(output)
    results = write_results_index(output, plan)
    raw_rows = []
    curve_rows = []
    for job in plan["jobs"]:
        result = next((item for item in results if item["job_id"] == job["job_id"]), None)
        benchmark = _lookup(plan, "benchmarks", str(job["benchmark_id"]))
        reducer = _lookup(plan, "reducers", str(job["reducer_id"]))
        row = {
            "order": job["order"], "job_id": job["job_id"], "wave": job["wave"],
            "case_id": job["benchmark_id"], "family": benchmark["family"],
            "theory": benchmark["theory"], "predicate_mode": benchmark["predicate_mode"],
            "reducer_id": job["reducer_id"], "repeat": job["repeat"],
            "status": result.get("status", "missing") if result else "missing",
            "verified": bool(result and result.get("verified")),
            "evidence_ok": bool(result and result.get("evidence_ok")),
            "predicate_calls": result.get("predicate_calls") if result else None,
            "accepted_moves": result.get("accepted_moves") if result else None,
            "trial_wall_sec": result.get("trial_wall_sec") if result else None,
            **_quality_columns(result.get("input_quality") if result else None, "input"),
            **_quality_columns(result.get("output_quality") if result else None, "output"),
            "evidence_warnings": json.dumps(result.get("evidence_warnings", []), separators=(",", ":")) if result else "[]",
        }
        raw_rows.append(row)
        if result:
            attempt_dir = Path(str(result["attempt_dir"]))
            trajectory = trajectory_for_attempt(attempt_dir, benchmark, reducer, allow_partial=False)
            for point in trajectory["points"]:
                curve_rows.append({
                    "job_id": job["job_id"], "case_id": job["benchmark_id"],
                    "reducer_id": job["reducer_id"], "repeat": job["repeat"],
                    **point,
                })

    by_reducer = {}
    for reducer in plan["reducers"]:
        selected = [item for item in results if item["reducer_id"] == reducer["id"]]
        formal = [item for item in selected if item.get("verified") and item.get("evidence_ok")]
        by_reducer[str(reducer["id"])] = {
            "label": reducer["label"], "planned": len(plan["benchmarks"]) * int(plan["repeats"]),
            "completed": len(selected), "verified": sum(bool(item.get("verified")) for item in selected),
            "evidence_ok": sum(bool(item.get("evidence_ok")) for item in selected),
            "statuses": dict(Counter(str(item["status"]) for item in selected)),
            "median_predicate_calls": _median(item.get("predicate_calls") for item in formal),
            "median_accepted_moves": _median(item.get("accepted_moves") for item in formal),
            "median_trial_wall_sec": _median(item.get("trial_wall_sec") for item in formal),
            "median_expression_ratio": _median(
                item["output_quality"]["expression_count"] / item["input_quality"]["expression_count"]
                for item in formal if item.get("input_quality") and item["input_quality"]["expression_count"]
            ),
            "median_node_ratio": _median(
                item["output_quality"]["node_count"] / item["input_quality"]["node_count"]
                for item in formal if item.get("input_quality") and item["input_quality"]["node_count"]
            ),
            "median_byte_ratio": _median(
                item["output_quality"]["byte_count"] / item["input_quality"]["byte_count"]
                for item in formal if item.get("input_quality") and item["input_quality"]["byte_count"]
            ),
        }

    comparisons = comparison_rows(plan, results)
    progress = {}
    try:
        progress = _mapping(_read_json(output / "progress.json", "progress"), "progress")
    except ReductionError:
        pass
    summary = {
        "schema_version": plan["schema_version"], "format": plan["format"],
        "study_id": plan["study_id"], "generated_at": _utc_now(),
        "status": progress.get("status", "unknown"),
        "total_jobs": len(plan["jobs"]), "completed_jobs": len(results),
        "verified_jobs": sum(bool(item.get("verified")) for item in results),
        "evidence_ok_jobs": sum(bool(item.get("evidence_ok")) for item in results),
        "benchmarks": len(plan["benchmarks"]), "repeats": plan["repeats"],
        "outer_jobs": plan["execution"]["outer_jobs"], "limits": plan["limits"],
        "selection": plan.get("selection", {}),
        "reducers": by_reducer, "comparisons": comparisons,
    }
    return summary, raw_rows, curve_rows


def _tsv(rows: Sequence[Mapping[str, object]]) -> bytes:
    if not rows:
        return b""
    from io import StringIO
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]), delimiter="\t", extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def report(output: Path) -> dict[str, object]:
    output = output.expanduser().resolve()
    summary, raw_rows, curve_rows = build_report(output)
    report_dir = output / "report"
    report_dir.mkdir(parents=True, exist_ok=True)
    _atomic_write(report_dir / "raw.tsv", _tsv(raw_rows))
    _atomic_write(report_dir / "curves.tsv", _tsv(curve_rows))
    _write_json(report_dir / "summary.json", summary)
    plan = load_plan(output)
    _write_json(report_dir / "report.complete.json", {
        "schema_version": plan["schema_version"],
        "format": plan["format"],
        "plan_sha256": plan.get("plan_sha256"),
        "results_sha256": _sha256_path(output / "results.tsv") if (output / "results.tsv").is_file() else None,
        "summary_sha256": _sha256_path(report_dir / "summary.json"),
        "raw_sha256": _sha256_path(report_dir / "raw.tsv"),
        "curves_sha256": _sha256_path(report_dir / "curves.tsv"),
    })
    return summary


def export_xlsx(output: Path, destination: Path | None = None) -> Path:
    from openpyxl import Workbook

    output = output.expanduser().resolve()
    summary, raw_rows, curve_rows = build_report(output)
    destination = (destination or output / "report" / f"{output.name}.xlsx").expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("summary")
    summary_sheet.append(["key", "value"])
    for key, value in summary.items():
        summary_sheet.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    jobs_sheet = workbook.create_sheet("jobs")
    if raw_rows:
        jobs_sheet.append(list(raw_rows[0]))
        for row in raw_rows:
            jobs_sheet.append([row.get(key) for key in raw_rows[0]])
    case_sheet = workbook.create_sheet("case-reducer")
    cases = case_rows(output)
    if cases:
        case_sheet.append(["case_id", "family", "theory", "predicate_mode", "planned", "completed", "verified", "evidence_ok", "statuses", "by_reducer"])
        for row in cases:
            case_sheet.append([
                row.get("case_id"), row.get("family"), row.get("theory"),
                row.get("predicate_mode"), row.get("planned"), row.get("completed"),
                row.get("verified"), row.get("evidence_ok"),
                json.dumps(row.get("statuses", {}), ensure_ascii=False),
                json.dumps(row.get("by_reducer", {}), ensure_ascii=False),
            ])
    reducer_sheet = workbook.create_sheet("reducers")
    reducer_sheet.append(["reducer_id", "label", "planned", "completed", "verified", "evidence_ok", "predicate_calls", "accepted_moves", "statuses"])
    for reducer_id, value in summary.get("reducers", {}).items():
        reducer_sheet.append([
            reducer_id, value.get("label"), value.get("planned"),
            value.get("completed"), value.get("verified"), value.get("evidence_ok"),
            value.get("median_predicate_calls", value.get("predicate_calls")),
            value.get("median_accepted_moves", value.get("accepted_moves")),
            json.dumps(value.get("statuses", {}), ensure_ascii=False),
        ])
    curves_sheet = workbook.create_sheet("trajectory")
    if curve_rows:
        curves_sheet.append(list(curve_rows[0]))
        for row in curve_rows:
            curves_sheet.append([row.get(key) for key in curve_rows[0]])
    evidence_sheet = workbook.create_sheet("evidence-health")
    evidence_sheet.append(["job_id", "case_id", "reducer_id", "verified", "evidence_ok", "warnings"])
    for row in raw_rows:
        evidence_sheet.append([
            row["job_id"], row["case_id"], row["reducer_id"],
            row["verified"], row["evidence_ok"], row["evidence_warnings"],
        ])
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]
    workbook.save(destination)
    return destination


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="action", required=True)
    prepare_parser = commands.add_parser("prepare", help="validate a v3 study and freeze its job plan")
    prepare_parser.add_argument("study", type=Path)
    prepare_parser.add_argument("--output", required=True, type=Path)
    prepare_parser.add_argument(
        "--reducers", nargs="+", required=True, metavar="ID",
        help="concrete reducer IDs from the project smtbatch.toml",
    )
    prepare_parser.add_argument(
        "--timeout", type=float, default=None,
        help="per-case reducer wall timeout in seconds (default: study value)",
    )
    prepare_parser.add_argument(
        "--jobs", type=int, default=None,
        help="outer parallel jobs (default: study value)",
    )
    run_parser = commands.add_parser("run", help="execute or resume a prepared plan with frozen outer jobs")
    run_parser.add_argument("output", type=Path)
    status_parser = commands.add_parser("status", help="show completion and evidence counts")
    status_parser.add_argument("output", type=Path)
    report_parser = commands.add_parser("report", help="build reduction summaries and trajectories")
    report_parser.add_argument("output", type=Path)
    report_parser.add_argument("--xlsx", action="store_true", help="also write the reduction workbook")
    stop_parser = commands.add_parser("stop", help="request a graceful or immediate stop")
    stop_parser.add_argument("output", type=Path)
    stop_parser.add_argument("--immediate", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.action == "prepare":
            plan = prepare(
                args.study, args.output, reducers=args.reducers,
                timeout_seconds=args.timeout, outer_jobs=args.jobs,
            )
            print(
                f"[reduce] prepared {len(plan['benchmarks'])} benchmarks x "
                f"{len(plan['reducers'])} reducers x {plan['repeats']} repeats "
                f"({len(plan['jobs'])} jobs, outer_jobs={plan['execution']['outer_jobs']})"
            )
            return 0
        if args.action == "run":
            results = run(args.output)
            state = status(args.output)
            print(f"[reduce] {state['status']} completed={len(results)}/{state['total_jobs']}")
            return 0 if state["status"] == "complete" else 130
        if args.action == "status":
            print(json.dumps(status(args.output), ensure_ascii=False, sort_keys=True))
            return 0
        if args.action == "report":
            value = report(args.output)
            if args.xlsx:
                path = export_xlsx(args.output)
                print(f"[reduce] workbook={path}")
            print(f"[reduce] report jobs={value['completed_jobs']} -> {args.output / 'report'}")
            return 0
        request = request_stop(args.output, "immediate" if args.immediate else "graceful")
        print(f"[reduce] requested {request['mode']} stop")
        return 0
    except (OSError, ReductionError, subprocess.SubprocessError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
